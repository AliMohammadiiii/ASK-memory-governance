#!/usr/bin/env python3
"""Build the editable English thesis review draft from the Markdown source."""

from __future__ import annotations

import argparse
import re
from collections.abc import Iterable
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
THESIS_DIR = ROOT / "thesis"
SOURCE = THESIS_DIR / "MASTERS_THESIS_KSA_MGP6.md"
BIB = THESIS_DIR / "references.bib"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT = "F4F6F9"
LIGHTER = "F8FAFC"
BORDER = "C9D2DD"
WHITE = "FFFFFF"
BLACK = "000000"
CITATION_LABELS: dict[str, str] = {}


def set_run_font(run, name: str, size: float | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    assert sum(widths) == CONTENT_WIDTH_DXA
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_cant_split(row)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_border(paragraph, *, bottom: str | None = None, left: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for side, color in (("bottom", bottom), ("left", left)):
        if not color:
            continue
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "10" if side == "left" else "6")
        node.set(qn("w:space"), "5")
        node.set(qn("w:color"), color)
        borders.append(node)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.333
    pf.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
        "Heading 4": (11, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False

    for name in ("KSA Code", "KSA Quote", "KSA Bibliography"):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    code = styles["KSA Code"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(3)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.keep_together = True

    quote = styles["KSA Quote"]
    quote.font.name = "Calibri"
    quote._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    quote._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    quote.font.size = Pt(10.5)
    quote.font.color.rgb = RGBColor.from_string(NAVY)
    quote.paragraph_format.left_indent = Inches(0.22)
    quote.paragraph_format.right_indent = Inches(0.12)
    quote.paragraph_format.space_before = Pt(6)
    quote.paragraph_format.space_after = Pt(8)
    quote.paragraph_format.line_spacing = 1.2

    bibliography = styles["KSA Bibliography"]
    bibliography.font.name = "Calibri"
    bibliography._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    bibliography._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    bibliography.font.size = Pt(9.5)
    bibliography.paragraph_format.left_indent = Inches(0.25)
    bibliography.paragraph_format.first_line_indent = Inches(-0.25)
    bibliography.paragraph_format.space_after = Pt(4)
    bibliography.paragraph_format.line_spacing = 1.15


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("KSA Memory Governance")
    set_run_font(left, "Calibri", 8.5, MUTED)
    right = p.add_run("\tEnglish Review Draft")
    set_run_font(right, "Calibri", 8.5, MUTED)
    set_paragraph_border(p, bottom=BORDER)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    run = p.add_run("Page ")
    set_run_font(run, "Calibri", 8.5, MUTED)
    add_field(p, "PAGE")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(70)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    r = kicker.add_run("MASTER'S THESIS - ENGLISH REVIEW DRAFT")
    set_run_font(r, "Calibri", 10, BLUE)
    r.bold = True
    r.font.small_caps = True

    p = doc.add_paragraph(
        "Auditable Governance of Long-Term Memory in Language Agents", style="Title"
    )
    p.paragraph_format.keep_with_next = True

    p = doc.add_paragraph(
        "The KSA Controller (Keep-Supersede-Abstain), the MGP-6 v2 Diagnostic, and Evidence-Bounded Evaluation",
        style="Subtitle",
    )
    p.paragraph_format.space_after = Pt(34)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(28)
    set_paragraph_border(rule, bottom=BLUE)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(18)
    r = metadata.add_run(
        "Review state: 15 July 2026\n"
        "Evidence basis: ledger v1.1 + construct-validity audit\n"
        "Editable research draft"
    )
    set_run_font(r, "Calibri", 11, MUTED)

    note = doc.add_paragraph(style="KSA Quote")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.left_indent = Inches(0.55)
    note.paragraph_format.right_indent = Inches(0.55)
    r = note.add_run(
        "The official University of Tehran template and verified administrative fields were not supplied. "
        "This document is a review copy, not an institution-ready submission."
    )
    r.bold = True
    shade_paragraph(note, LIGHT)
    set_paragraph_border(note, left=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Administrative fields intentionally remain external")
    set_run_font(r, "Calibri", 9.5, MUTED)

    doc.add_page_break()
    front = doc.add_paragraph("Front Matter", style="Heading 1")
    front.paragraph_format.page_break_before = False


def add_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    existing_abs = [
        int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))
    ]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs, default=-1) + 1
    next_num = max(existing_num, default=0) + 1

    def make(marker: str, fmt: str, abs_id: int, num_id: int) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), marker)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "280")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abs_id))
        num.append(ref)
        numbering.append(num)

    make("•", "bullet", next_abs, next_num)
    make("%1.", "decimal", next_abs + 1, next_num + 1)
    return next_num, next_num + 1


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208


def clone_numbering_instance(doc: Document, template_num_id: int) -> int:
    numbering = doc.part.numbering_part.element
    template = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == template_num_id
    )
    abstract_ref = template.find(qn("w:abstractNumId"))
    if abstract_ref is None:
        raise ValueError(f"Numbering template {template_num_id} has no abstract reference")
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_ref.get(qn("w:val")))
    num.append(ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))")


def format_math_markup(text: str) -> str:
    subscript = str.maketrans("0123456789itn", "₀₁₂₃₄₅₆₇₈₉ᵢₜₙ")

    def convert_inline_math(match: re.Match[str]) -> str:
        value = match.group(1)
        value = re.sub(
            r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"\1/\2", value
        )
        value = re.sub(
            r"\\sum_\{?([A-Za-z0-9]+)\}?",
            lambda item: "Σ" + item.group(1).translate(subscript),
            value,
        )
        value = re.sub(
            r"_\{?([0-9itn])\}?",
            lambda item: item.group(1).translate(subscript),
            value,
        )
        return value

    text = re.sub(r"\\\((.*?)\\\)", convert_inline_math, text)
    text = re.sub(r"\\text\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\mathbb\{([^{}]*)\}", r"\1", text)
    text = text.replace(r"\pi", "π").replace(r"\theta", "θ")
    text = text.replace(r"\in", "∈").replace(r"\Pr", "Pr")
    text = text.replace(r"\mid", "|")
    text = text.replace(r"\hat y", "ŷ").replace(r"\hat", "")
    text = text.replace(r"\{", "{").replace(r"\}", "}")
    text = text.replace(r"\,", " ")
    return text


def format_citations(text: str) -> str:
    def replace(match: re.Match[str]) -> str:
        keys = re.findall(r"@([A-Za-z0-9_:-]+)", match.group(1))
        if not keys or any(key not in CITATION_LABELS for key in keys):
            return match.group(0)
        return "(" + "; ".join(CITATION_LABELS[key] for key in keys) + ")"

    return re.sub(r"\[((?:@[A-Za-z0-9_:-]+(?:;\s*)?)+)\]", replace, text)


def add_inline(paragraph, text: str, *, base_size: float | None = None) -> None:
    text = format_citations(format_math_markup(text))
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if base_size:
                set_run_font(run, "Calibri", base_size)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", (base_size or 10.5) - 0.5, DARK_BLUE)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            if base_size:
                set_run_font(run, "Calibri", base_size)
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            if base_size:
                set_run_font(run, "Calibri", base_size)
            run.italic = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if base_size:
            set_run_font(run, "Calibri", base_size)


def clean_cell_text(value: str) -> str:
    value = value.strip()
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"(?<!\*)\*(.*?)\*(?!\*)", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    return value.replace("<br>", "\n")


def column_widths(rows: list[list[str]]) -> list[int]:
    cols = len(rows[0])
    headers = [clean_cell_text(cell).casefold() for cell in rows[0]]
    if headers == ["research question", "main chapter(s)", "primary artifact"]:
        return [1900, 1900, 5560]
    lengths = []
    for idx in range(cols):
        max_len = max(len(clean_cell_text(row[idx])) for row in rows)
        header_len = len(clean_cell_text(rows[0][idx]))
        lengths.append(max(5, min(38, max(max_len, header_len * 2))))
    minimum = 760 if cols >= 6 else 900 if cols >= 5 else 1100 if cols >= 4 else 1250
    available = CONTENT_WIDTH_DXA - minimum * cols
    total = sum(lengths)
    widths = [minimum + int(available * length / total) for length in lengths]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    cols = len(rows[0])
    rows = [row[:cols] + [""] * max(0, cols - len(row)) for row in rows]
    widths = column_widths(rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    font_size = 8.1 if cols >= 7 else 8.5 if cols == 6 else 9.0 if cols >= 4 else 9.5

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if c_idx > 0 and len(value) < 28
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(p, clean_cell_text(value), base_size=font_size)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHTER)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_picture(doc: Document, alt: str, path_text: str) -> None:
    path = (THESIS_DIR / path_text).resolve()
    if not path.exists():
        p = doc.add_paragraph(style="KSA Quote")
        add_inline(p, f"Figure unavailable in review build: {alt} ({path_text})")
        shade_paragraph(p, LIGHT)
        set_paragraph_border(p, left=BLUE)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    drawing = run._r.find(qn("w:drawing"))
    if drawing is not None:
        doc_pr = drawing.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("descr", alt)
            doc_pr.set("title", alt)
    caption = doc.add_paragraph(alt, style="Caption")
    caption.paragraph_format.keep_with_next = False


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw = lines[idx].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def add_code_block(doc: Document, code_lines: Iterable[str]) -> None:
    text = "\n".join(code_lines).rstrip()
    p = doc.add_paragraph(style="KSA Code")
    r = p.add_run(text)
    set_run_font(r, "Consolas", 8.5, NAVY)
    shade_paragraph(p, LIGHT)
    set_paragraph_border(p, left=BLUE)


def add_markdown(doc: Document, source: str, bullet_num_id: int, decimal_num_id: int) -> set[str]:
    citations: set[str] = set()
    for group in re.findall(r"\[([^\]]*@[^\]]+)\]", source):
        citations.update(re.findall(r"@([A-Za-z0-9_:-]+)", group))
    citations.discard("key")  # literal syntax example in the References note
    lines = source.splitlines()
    index = 3  # cover already carries the first H1/H2 title block
    in_code = False
    code_lines: list[str] = []
    in_math = False
    math_lines: list[str] = []
    list_kind: str | None = None
    active_bullet_num = bullet_num_id
    active_decimal_num = decimal_num_id

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        bullet = re.match(r"^-\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if not bullet and not numbered:
            list_kind = None

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped == "\\[":
            in_math = True
            math_lines = []
            index += 1
            continue
        if stripped == "\\]" and in_math:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(format_math_markup(" ".join(math_lines)))
            set_run_font(r, "Cambria Math", 10.5, NAVY)
            in_math = False
            index += 1
            continue
        if in_math:
            math_lines.append(stripped)
            index += 1
            continue

        if not stripped:
            index += 1
            continue
        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            set_paragraph_border(p, bottom=BORDER)
            index += 1
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            add_picture(doc, image_match.group(1).rstrip("."), image_match.group(2))
            index += 1
            continue

        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and lines[index + 1].strip().startswith("|")
        ):
            rows, index = parse_table(lines, index)
            if rows:
                add_markdown_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            p = doc.add_paragraph(style=f"Heading {level}")
            if level == 1 and (text.startswith("1. ") or text == "References"):
                p.paragraph_format.page_break_before = True
            add_inline(p, text)
            index += 1
            continue

        if bullet:
            if list_kind != "bullet":
                active_bullet_num = clone_numbering_instance(doc, bullet_num_id)
            list_kind = "bullet"
            p = doc.add_paragraph()
            apply_numbering(p, active_bullet_num)
            add_inline(p, bullet.group(1))
            index += 1
            continue

        if numbered:
            if list_kind != "decimal":
                active_decimal_num = clone_numbering_instance(doc, decimal_num_id)
            list_kind = "decimal"
            p = doc.add_paragraph()
            apply_numbering(p, active_decimal_num)
            add_inline(p, numbered.group(1))
            index += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph(style="KSA Quote")
            add_inline(p, stripped.lstrip("> "))
            shade_paragraph(p, LIGHT)
            set_paragraph_border(p, left=BLUE)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (
                not nxt
                or nxt.startswith(("#", "|", "- ", ">", "```", "![", "\\["))
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            paragraph_lines.append(nxt)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))

    return citations


def parse_bibtex(text: str) -> dict[str, dict[str, str]]:
    entries: dict[str, dict[str, str]] = {}
    pos = 0
    while True:
        match = re.search(r"@(\w+)\s*\{\s*([^,]+),", text[pos:])
        if not match:
            break
        body_start = pos + match.end()
        depth = 1
        idx = body_start
        while idx < len(text) and depth:
            if text[idx] == "{":
                depth += 1
            elif text[idx] == "}":
                depth -= 1
            idx += 1
        body = text[body_start : idx - 1]
        fields: dict[str, str] = {"entry_type": match.group(1), "key": match.group(2).strip()}
        field_pattern = re.compile(r"(?ms)^\s*(\w+)\s*=\s*\{(.*?)\}\s*,?\s*$")
        for field in field_pattern.finditer(body):
            fields[field.group(1).lower()] = field.group(2).replace("\n", " ").strip()
        entries[fields["key"]] = fields
        pos = idx
    return entries


def latex_to_text(value: str) -> str:
    replacements = {
        r"{\'o}": "ó",
        r"{\'e}": "é",
        r"{\'a}": "á",
        r"{\'U}": "Ú",
        r"{\`e}": "è",
        r"{\"a}": "ä",
        r"{\"u}": "ü",
        r"{\c{c}}": "ç",
        r"{\v{s}}": "š",
        r"{\v{z}}": "ž",
        r"{\o}": "ø",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    return value.replace("{", "").replace("}", "").replace("--", "-")


def build_citation_labels(entries: dict[str, dict[str, str]]) -> dict[str, str]:
    labels: dict[str, str] = {}
    for key, entry in entries.items():
        author_raw = entry.get("author", "").strip()
        year = latex_to_text(entry.get("year", "n.d."))
        if author_raw:
            authors = [latex_to_text(part.strip()) for part in author_raw.split(" and ")]
            first = authors[0]
            if author_raw.startswith("{{"):
                lead = first
            elif "," in first:
                lead = first.split(",", 1)[0].strip()
            else:
                lead = first.split()[-1]
            if len(authors) == 2 and authors[1].casefold() != "others":
                second = (
                    authors[1].split(",", 1)[0].strip()
                    if "," in authors[1]
                    else authors[1].split()[-1]
                )
                lead = f"{lead} and {second}"
            elif len(authors) > 2 or authors[-1].casefold() == "others":
                lead = f"{lead} et al."
        else:
            title = latex_to_text(entry.get("title", key))
            lead = title.split(":", 1)[0]
        labels[key] = f"{lead}, {year}"
    return labels


def add_bibliography(doc: Document, citations: set[str]) -> None:
    entries = parse_bibtex(BIB.read_text(encoding="utf-8"))
    missing = sorted(citations - entries.keys())
    if missing:
        raise ValueError(f"Missing bibliography entries: {', '.join(missing)}")
    heading = doc.add_paragraph("Compiled reference list", style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    for key in sorted(
        citations,
        key=lambda item: (entries[item].get("author", ""), entries[item].get("year", ""), item),
    ):
        entry = entries[key]
        author = latex_to_text(entry.get("author", "Unknown author")).replace(" and ", "; ")
        year = latex_to_text(entry.get("year", "n.d."))
        title = latex_to_text(entry.get("title", key))
        venue = latex_to_text(
            entry.get("journal")
            or entry.get("booktitle")
            or entry.get("howpublished")
            or entry.get("publisher")
            or ""
        )
        doi = latex_to_text(entry.get("doi", ""))
        parts = [f"{author} ({year}). {title}."]
        if venue:
            parts.append(venue + ".")
        if doi:
            parts.append("DOI: " + doi + ".")
        p = doc.add_paragraph(style="KSA Bibliography")
        add_inline(p, " ".join(parts))


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Auditable Governance of Long-Term Memory in Language Agents"
    props.subject = "Master's thesis English review draft"
    props.keywords = "memory governance; KSA; MGP-6; long-term memory; language agents"
    props.comments = "Review draft. Official University of Tehran template and administrative fields remain external."
    props.author = ""
    props.last_modified_by = ""


def build(output: Path) -> None:
    global CITATION_LABELS
    source = SOURCE.read_text(encoding="utf-8")
    CITATION_LABELS = build_citation_labels(parse_bibtex(BIB.read_text(encoding="utf-8")))
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    bullet_num_id, decimal_num_id = add_numbering(doc)
    add_cover(doc)
    citations = add_markdown(doc, source, bullet_num_id, decimal_num_id)
    add_bibliography(doc, citations)
    set_core_properties(doc)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=THESIS_DIR / "compiled" / "KSA_Masters_Thesis_English_Review_Draft.docx",
    )
    args = parser.parse_args()
    build(args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
