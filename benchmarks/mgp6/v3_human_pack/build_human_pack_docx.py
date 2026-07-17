#!/usr/bin/env python3
"""Build polished RTL DOCX files from the Persian MGP-6 v3 role manuals."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DIST = ROOT / "dist"
FONT = "IRANYekanX"
FONT_LATIN = "Arial"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEXT = "1F2937"
MUTED = "64748B"
PALE_BLUE = "E8EEF5"
PALE_YELLOW = "FFF4CC"
PALE_GRAY = "F6F8FA"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360

FILES = {
    "00_COORDINATOR_GUIDE_FA.md": (
        "MGP6_V3_COORDINATOR_GUIDE_FA.docx",
        "راهنمای هماهنگ‌کننده — محرمانه",
        "COORDINATOR ONLY",
    ),
    "01_WRITER_INSTRUCTIONS_FA.md": (
        "MGP6_V3_WRITER_INSTRUCTIONS_FA.docx",
        "بسته قابل ارسال برای نویسنده داده",
        "DATA WRITER",
    ),
    "02_ANNOTATOR_INSTRUCTIONS_FA.md": (
        "MGP6_V3_ANNOTATOR_INSTRUCTIONS_FA.docx",
        "بسته قابل ارسال برای برچسب‌زن مستقل",
        "ANNOTATOR",
    ),
    "03_ADJUDICATOR_INSTRUCTIONS_FA.md": (
        "MGP6_V3_ADJUDICATOR_INSTRUCTIONS_FA.docx",
        "بسته قابل ارسال برای داور چهارم",
        "ADJUDICATOR",
    ),
    "04_PARTICIPANT_CONSENT_FA.md": (
        "MGP6_V3_PARTICIPANT_CONSENT_FA.docx",
        "فرم اطلاع‌رسانی، رضایت و بازتوزیع",
        "CONSENT FORM",
    ),
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="CBD5E1", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tbl_pr.append(bidi)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, *, size=None, bold=None, color=None, latin=False) -> None:
    run.font.name = FONT_LATIN if latin else FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN if latin else FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN if latin else FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "fa-IR")
    lang.set(qn("w:bidi"), "fa-IR")
    if not latin:
        rtl = run._element.get_or_add_rPr().find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl")
            run._element.get_or_add_rPr().append(rtl)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_rtl(paragraph, *, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    paragraph.alignment = align


def set_paragraph_shading(paragraph, fill: str, border_color: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border_color:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        right = OxmlElement("w:right")
        right.set(qn("w:val"), "single")
        right.set(qn("w:sz"), "18")
        right.set(qn("w:space"), "8")
        right.set(qn("w:color"), border_color)
        p_bdr.append(right)


INLINE_RE = re.compile(r"(\*\*.*?\*\*|`.*?`|\[[^\]]+\])")


def add_inline(paragraph, text: str, *, default_bold=False, color=TEXT, size=11) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_font(run, size=size, bold=default_bold, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=max(size - 0.5, 9), color=DARK_BLUE, latin=True)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_GRAY)
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(token)
            set_font(run, size=size, bold=True, color="7C5C00")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_YELLOW)
            run._element.get_or_add_rPr().append(shd)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run, size=size, bold=default_bold, color=color)


class NumberingManager:
    def __init__(self, document: Document):
        self.numbering = document.part.numbering_part.element
        self.abstract_bullet = self._new_abstract(bullet=True)
        self.abstract_decimal = self._new_abstract(bullet=False)

    def _next_id(self, tag: str, attr: str) -> int:
        values = []
        for node in self.numbering.findall(qn(tag)):
            value = node.get(qn(attr))
            if value is not None:
                values.append(int(value))
        return max(values, default=-1) + 1

    def _new_abstract(self, *, bullet: bool) -> int:
        abstract_id = self._next_id("w:abstractNum", "w:abstractNumId")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:start"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        lvl.append(p_pr)
        abstract.append(lvl)
        self.numbering.append(abstract)
        return abstract_id

    def new_num(self, *, bullet: bool) -> int:
        num_id = self._next_id("w:num", "w:numId")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(self.abstract_bullet if bullet else self.abstract_decimal))
        num.append(abstract)
        lvl_override = OxmlElement("w:lvlOverride")
        lvl_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        lvl_override.append(start_override)
        num.append(lvl_override)
        self.numbering.append(num)
        return num_id

    @staticmethod
    def apply(paragraph, num_id: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_el)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    result_text = OxmlElement("w:t")
    result_text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(result_text)
    run._r.append(fld_char_end)
    set_font(run, size=8.5, color=MUTED, latin=True)


def configure_document(document: Document, title: str, role_tag: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style = styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:cs"), FONT)
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    header = section.header
    hp = header.paragraphs[0]
    set_rtl(hp, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_inline(hp, f"MGP-6 v3  |  {role_tag}", default_bold=True, color=MUTED, size=8.5)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("MGP-6 v3  •  ")
    set_font(run, size=8.5, color=MUTED, latin=True)
    add_page_field(fp)

    document.core_properties.title = title
    document.core_properties.subject = "MGP-6 v3 independent human operations pack"
    document.core_properties.author = "MGP-6 v3 research coordination"
    document.core_properties.keywords = "MGP-6, memory governance, annotation, adjudication, Persian"


def add_cover(document: Document, title: str, subtitle: str, role_tag: str, intro_lines: list[tuple[str, str]]) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(68)
    set_rtl(spacer)

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(16)
    run = rule.add_run("━" * 22)
    set_font(run, size=9, color=BLUE)

    p = document.add_paragraph(style="Title")
    set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_inline(p, title, default_bold=True, color=DARK_BLUE, size=24)
    p.paragraph_format.space_after = Pt(12)

    p = document.add_paragraph()
    set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_inline(p, subtitle, default_bold=True, color=BLUE, size=11.5)
    p.paragraph_format.space_after = Pt(8)

    p = document.add_paragraph()
    set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_inline(p, role_tag, default_bold=True, color=WHITE, size=9)
    set_paragraph_shading(p, BLUE)
    p.paragraph_format.left_indent = Inches(1.9)
    p.paragraph_format.right_indent = Inches(1.9)
    p.paragraph_format.space_after = Pt(18)

    for kind, text in intro_lines:
        p = document.add_paragraph()
        set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, text, color=TEXT if kind == "plain" else DARK_BLUE, size=10.5)
        if kind == "quote":
            set_paragraph_shading(p, PALE_BLUE, BLUE)
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.right_indent = Inches(0.45)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    add_inline(p, "نسخه ۳.۰  •  ۱۴۰۵/۰۴/۲۴", color=MUTED, size=9)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_heading(document: Document, level: int, text: str) -> None:
    p = document.add_paragraph(style=f"Heading {min(level, 3)}")
    set_rtl(p)
    add_inline(p, text, default_bold=True, color=BLUE if level < 3 else DARK_BLUE, size={1: 16, 2: 13, 3: 12}.get(level, 11))


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    set_rtl(p)
    add_inline(p, text)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.25


def add_quote(document: Document, text: str) -> None:
    p = document.add_paragraph()
    set_rtl(p)
    add_inline(p, text, default_bold=True, color=DARK_BLUE, size=10.5)
    set_paragraph_shading(p, PALE_BLUE, BLUE)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)


def add_list_item(document: Document, text: str, numbering: NumberingManager, num_id: int, *, checkbox=False) -> None:
    p = document.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if checkbox:
        text = "☐ " + text
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
    else:
        NumberingManager.apply(p, num_id)
    add_inline(p, text)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    if cols == 2:
        widths = [int(TABLE_WIDTH_DXA * 0.35), TABLE_WIDTH_DXA - int(TABLE_WIDTH_DXA * 0.35)]
    else:
        base = TABLE_WIDTH_DXA // cols
        widths = [base] * (cols - 1) + [TABLE_WIDTH_DXA - base * (cols - 1)]
    set_table_width(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(normalized):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, PALE_BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, default_bold=(r_idx == 0), size=9.5)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def is_table_separator(line: str) -> bool:
    parts = [part.strip() for part in line.strip().strip("|").split("|")]
    return bool(parts) and all(re.fullmatch(r":?-{3,}:?", part) for part in parts)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if is_table_separator(line):
            continue
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    return rows


def render_markdown(document: Document, markdown: str, title: str, subtitle: str, role_tag: str) -> None:
    lines = markdown.splitlines()
    first_title = title
    intro: list[tuple[str, str]] = []
    idx = 0
    if lines and lines[0].startswith("# "):
        first_title = lines[0][2:].strip()
        idx = 1
    while idx < len(lines):
        line = lines[idx].strip()
        if line.startswith("## "):
            break
        if line and line != "---":
            if line.startswith("> "):
                intro.append(("quote", line[2:].strip()))
            else:
                intro.append(("plain", line))
        idx += 1
    add_cover(document, first_title, subtitle, role_tag, intro)

    numbering = NumberingManager(document)
    active_list_kind = None
    active_num_id = None
    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()
        if not line or line == "---":
            active_list_kind = None
            active_num_id = None
            idx += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                table_lines.append(lines[idx].strip())
                idx += 1
            add_table(document, parse_table(table_lines))
            active_list_kind = None
            active_num_id = None
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            add_heading(document, len(heading.group(1)), heading.group(2))
            active_list_kind = None
            active_num_id = None
            idx += 1
            continue
        if line.startswith("> "):
            add_quote(document, line[2:].strip())
            active_list_kind = None
            active_num_id = None
            idx += 1
            continue
        checkbox = re.match(r"^- \[[ xX]\]\s+(.+)$", line)
        bullet = re.match(r"^-\s+(.+)$", line)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if checkbox or bullet:
            if active_list_kind != "bullet":
                active_num_id = numbering.new_num(bullet=True)
                active_list_kind = "bullet"
            add_list_item(
                document,
                (checkbox or bullet).group(1),
                numbering,
                active_num_id,
                checkbox=bool(checkbox),
            )
            idx += 1
            continue
        if numbered:
            if active_list_kind != "number":
                active_num_id = numbering.new_num(bullet=False)
                active_list_kind = "number"
            add_list_item(document, numbered.group(1), numbering, active_num_id)
            idx += 1
            continue

        paragraph_lines = [line]
        idx += 1
        while idx < len(lines):
            nxt = lines[idx].strip()
            if (
                not nxt
                or nxt == "---"
                or nxt.startswith("#")
                or nxt.startswith("|")
                or nxt.startswith("> ")
                or re.match(r"^-\s+", nxt)
                or re.match(r"^\d+[.)]\s+", nxt)
            ):
                break
            paragraph_lines.append(nxt)
            idx += 1
        add_body(document, " ".join(paragraph_lines))
        active_list_kind = None
        active_num_id = None


def build_one(source_name: str, output_name: str, subtitle: str, role_tag: str) -> Path:
    source = ROOT / source_name
    document = Document()
    configure_document(document, source.stem, role_tag)
    render_markdown(document, source.read_text(encoding="utf-8"), source.stem, subtitle, role_tag)
    DIST.mkdir(parents=True, exist_ok=True)
    output = DIST / output_name
    document.save(output)
    return output


def main() -> None:
    outputs = []
    for source_name, (output_name, subtitle, role_tag) in FILES.items():
        outputs.append(build_one(source_name, output_name, subtitle, role_tag))
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
